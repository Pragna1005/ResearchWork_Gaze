"""Build the progress presentation (pptx) plus PNG previews of every slide.

    python presentation/make_charts.py      # charts first
    python presentation/build_deck.py       # -> presentation/Gaze_Confusion_Progress.pptx, preview/

The previews are drawn with PIL using the same boxes and fonts, so text overflow
can be checked without PowerPoint. Edit SPEAKER / content below and rebuild.
"""
import os
import re

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
CHARTS = os.path.join(HERE, "charts")
PREVIEW = os.path.join(HERE, "preview")
SPEAKER = "[Your name]"

W_IN, H_IN = 13.333, 7.5
PX = 120  # preview pixels per inch
INK, INK2, MUTED = "0b0b0b", "52514e", "8a8984"
BLUE, ORANGE, GREY = "2a78d6", "eb6834", "a3a29d"
BLUE_BG, ORANGE_BG, GREY_BG = "eaf2fc", "fdeee7", "f3f3f1"
FONT = "Calibri"
FONT_FILES = {False: r"C:\Windows\Fonts\calibri.ttf", True: r"C:\Windows\Fonts\calibrib.ttf"}


def rgb(h):
    return RGBColor.from_string(h)


def runs_of(text):
    """'plain **bold** plain' -> [(text, bold), ...]"""
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    return [(p[2:-2], True) if p.startswith("**") else (p, False) for p in parts if p]


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W_IN), Inches(H_IN)
        self.blank = self.prs.slide_layouts[6]
        self.previews = []
        self.overflows = []
        self.n = 0

    # ---------------------------------------------------------------- slides
    def new(self, title, notes, kicker=None):
        self.n += 1
        self.slide = self.prs.slides.add_slide(self.blank)
        self.img = Image.new("RGB", (int(W_IN * PX), int(H_IN * PX)), "white")
        self.draw = ImageDraw.Draw(self.img)
        self.previews.append(self.img)
        self.slide.notes_slide.notes_text_frame.text = notes.strip()
        self.rect(0, 0, 0.18, H_IN, BLUE)
        if kicker:
            self.text(0.7, 0.38, 11.5, 0.35, [{"t": kicker.upper(), "size": 12, "color": BLUE, "bold": True}])
        self.text(0.7, 0.68, 12.0, 1.15, [{"t": title, "size": 28, "bold": True}], name="title")
        self.text(0.7, 7.0, 9, 0.3, [{"t": "Gaze-based confusion detection · progress report", "size": 10, "color": MUTED}])
        self.text(12.0, 7.0, 0.7, 0.3, [{"t": str(self.n), "size": 10, "color": MUTED, "align": "r"}])

    # ---------------------------------------------------------------- shapes
    def rect(self, x, y, w, h, fill, line=None, rounded=False):
        shp = self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                          Inches(x), Inches(y), Inches(w), Inches(h))
        if rounded:
            shp.adjustments[0] = 0.08
        shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill)
        if line:
            shp.line.color.rgb = rgb(line); shp.line.width = Pt(1)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        box = [x * PX, y * PX, (x + w) * PX, (y + h) * PX]
        if rounded:
            self.draw.rounded_rectangle(box, radius=0.08 * min(w, h) * PX, fill="#" + fill,
                                        outline=("#" + line) if line else None)
        else:
            self.draw.rectangle(box, fill="#" + fill, outline=("#" + line) if line else None)
        return shp

    def arrow(self, x, y, w, h, fill=GREY):
        shp = self.slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
        shp.fill.solid(); shp.fill.fore_color.rgb = rgb(fill); shp.line.fill.background(); shp.shadow.inherit = False
        X, Y, Wp, Hp = x * PX, y * PX, w * PX, h * PX
        self.draw.polygon([(X, Y + Hp * .25), (X + Wp * .5, Y + Hp * .25), (X + Wp * .5, Y), (X + Wp, Y + Hp / 2),
                           (X + Wp * .5, Y + Hp), (X + Wp * .5, Y + Hp * .75), (X, Y + Hp * .75)], fill="#" + fill)

    def picture(self, path, x, y, w=None, h=None):
        im = Image.open(path)
        ar = im.height / im.width
        if w is None:
            w = h / ar
        h = w * ar
        self.slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
        self.img.paste(im.convert("RGB").resize((int(w * PX), int(h * PX)), Image.LANCZOS), (int(x * PX), int(y * PX)))
        return h

    # ---------------------------------------------------------------- text
    def text(self, x, y, w, h, paras, anchor="top", name=None, check=True):
        """paras: list of dicts: t, size, bold, color, bullet, align ('l'/'c'/'r'), after (pt)."""
        tb = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE}[anchor]
        layout = []
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            size = p.get("size", 18)
            para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[p.get("align", "l")]
            para.space_after = Pt(p.get("after", 6))
            para.line_spacing = 1.0
            for txt, bold in runs_of(p["t"]):
                r = para.add_run(); r.text = txt
                r.font.name = FONT; r.font.size = Pt(size)
                r.font.bold = p.get("bold", False) or bold
                r.font.color.rgb = rgb(p.get("color", INK))
            if p.get("bullet"):
                ind = Inches(0.26 if size >= 16 else 0.22)
                pPr = para._p.get_or_add_pPr()
                pPr.set("marL", str(Emu(ind))); pPr.set("indent", str(-Emu(ind)))
                bc = pPr.makeelement(qn("a:buClr"), {}); clr = bc.makeelement(qn("a:srgbClr"), {"val": p.get("bcolor", BLUE)})
                bc.append(clr); pPr.append(bc)
                bu = pPr.makeelement(qn("a:buChar"), {"char": "•"}); pPr.append(bu)
            layout.append(p)
        self._preview_text(x, y, w, h, layout, anchor, name, check)
        return tb

    def _preview_text(self, x, y, w, h, paras, anchor, name, check):
        lines_all = []  # (paragraph, [line of (word, bold)], indent_px, line_h, after_px)
        total = 0
        for p in paras:
            size_px = p.get("size", 18) * PX / 72
            fonts = {b: ImageFont.truetype(FONT_FILES[b], int(round(size_px))) for b in (False, True)}
            ind = (0.26 if p.get("size", 18) >= 16 else 0.22) * PX if p.get("bullet") else 0
            maxw = w * PX - ind
            words = []
            for txt, bold in runs_of(p["t"]):
                for tok in re.split(r"(\s+)", txt):
                    if tok:
                        words.append((tok, p.get("bold", False) or bold))
            lines, cur, curw = [], [], 0
            for tok, b in words:
                tw = fonts[b].getlength(tok)
                if tok.isspace():
                    if cur:
                        cur.append((tok, b)); curw += tw
                    continue
                if curw + tw > maxw and cur:
                    while cur and cur[-1][0].isspace():
                        curw -= fonts[cur[-1][1]].getlength(cur[-1][0]); cur.pop()
                    lines.append(cur); cur, curw = [], 0
                cur.append((tok, b)); curw += tw
            if cur:
                lines.append(cur)
            lh = size_px * 1.2
            after = p.get("after", 6) * PX / 72
            lines_all.append((p, lines, ind, lh, after, fonts))
            total += lh * len(lines) + after
        total -= lines_all[-1][4] if lines_all else 0
        if check and total > h * PX + 2:
            self.overflows.append(f"slide {self.n}: text box '{(name or paras[0]['t'][:40])}' needs "
                                  f"{total / PX:.2f} in, has {h:.2f} in")
        cy = y * PX + ((h * PX - total) / 2 if anchor == "middle" else 0)
        for p, lines, ind, lh, after, fonts in lines_all:
            color = "#" + p.get("color", INK)
            for li, line in enumerate(lines):
                lw = sum(fonts[b].getlength(t) for t, b in line)
                cx = x * PX + ind
                if p.get("align", "l") == "c":
                    cx = x * PX + (w * PX - lw) / 2
                elif p.get("align", "l") == "r":
                    cx = x * PX + w * PX - lw
                if p.get("bullet") and li == 0:
                    self.draw.text((x * PX, cy), "•", font=fonts[False], fill="#" + p.get("bcolor", BLUE))
                for t, b in line:
                    self.draw.text((cx, cy), t, font=fonts[b], fill=color)
                    cx += fonts[b].getlength(t)
                cy += lh
            cy += after

    # ---------------------------------------------------------------- building blocks
    def card(self, x, y, w, h, head, body, fill=GREY_BG, head_color=INK, body_size=15, head_size=17):
        self.rect(x, y, w, h, fill, rounded=True)
        paras = [{"t": head, "size": head_size, "bold": True, "color": head_color, "after": 4}]
        paras += [{"t": b, "size": body_size, "color": INK2, "after": 3} for b in ([body] if isinstance(body, str) else body)]
        self.text(x + 0.2, y + 0.16, w - 0.4, h - 0.3, paras)

    def stat(self, x, y, w, h, number, caption, color=BLUE, fill=GREY_BG):
        self.rect(x, y, w, h, fill, rounded=True)
        self.text(x + 0.2, y + 0.15, w - 0.4, 0.75, [{"t": number, "size": 34, "bold": True, "color": color}])
        self.text(x + 0.2, y + 0.95, w - 0.4, h - 1.05, [{"t": caption, "size": 14, "color": INK2}])

    def bullets(self, x, y, w, h, items, size=18, after=8):
        self.text(x, y, w, h, [{"t": t, "size": size, "bullet": True, "after": after} for t in items])

    def save(self, path):
        self.prs.save(path)
        os.makedirs(PREVIEW, exist_ok=True)
        for i, im in enumerate(self.previews, 1):
            im.save(os.path.join(PREVIEW, f"slide_{i:02d}.png"))


d = Deck()

# 1 ---------------------------------------------------------------- title
d.n += 1
d.slide = d.prs.slides.add_slide(d.blank)
d.img = Image.new("RGB", (int(W_IN * PX), int(H_IN * PX)), "white"); d.draw = ImageDraw.Draw(d.img); d.previews.append(d.img)
d.rect(0, 0, 0.18, H_IN, BLUE)
d.text(0.9, 1.7, 11.5, 0.4, [{"t": "PROGRESS REPORT", "size": 14, "bold": True, "color": BLUE}])
d.text(0.9, 2.15, 11.5, 1.8, [{"t": "Detecting reading confusion from eye movements", "size": 44, "bold": True}])
d.text(0.9, 3.95, 11.0, 1.0, [{"t": "Reproducing the pipeline, fixing what was broken, and an honest first result",
                               "size": 24, "color": INK2}])
d.rect(0.9, 5.2, 1.2, 0.05, BLUE)
d.text(0.9, 5.45, 11.0, 1.0, [{"t": SPEAKER, "size": 18, "bold": True, "after": 2},
                              {"t": "Ubiquitous & Systems Research Lab, IIT Jodhpur · October 2026", "size": 16, "color": INK2}])
d.slide.notes_slide.notes_text_frame.text = (
    "Good morning. In this talk I'll explain what I found in the project I inherited, what I fixed, the two experiments "
    "I ran, and what I propose next. Short version: the old model's accuracy was not valid, but after fixing the data "
    "and evaluating honestly, eye movements alone still detect confusion clearly above chance on texts the model has "
    "never seen.")

# 2 ---------------------------------------------------------------- overview
d.new("What I did, in five steps", kicker="Overview", notes="""
This is the whole talk on one slide. First I reproduced the existing system to make sure I understood it and that it
runs on our lab PC. Then I audited the data and the old model, and found why its accuracy cannot be trusted. Third, I
fixed bugs in how eye movements were measured. Fourth, I ran the two experiments the handover document asked for:
relabelling the data by time, and training a new model evaluated honestly. Finally, I fixed the recording script so the
new data we collect doesn't repeat the old problems. Everything is saved in git, so every change can be shown as a
before-and-after.""")
steps = [("1", "Reproduced", "Ran the whole Tobii pipeline on the lab PC; set up the webcam app"),
         ("2", "Audited", "Checked the old data and model — found why the old accuracy is not valid"),
         ("3", "Fixed measurements", "Regressions, saccade timing, fixations lost on blinks"),
         ("4", "Experiments", "Exp. 1: labels by time. Exp. 2: honest test on unseen texts, AUC 0.76"),
         ("5", "Fixed recording", "29 Hz gaze, one folder per session, quality check, protocol")]
cw, gap, x0 = 2.25, 0.18, 0.7
for i, (num, head, body) in enumerate(steps):
    x = x0 + i * (cw + gap)
    d.rect(x, 1.95, cw, 3.6, BLUE_BG if i in (3,) else GREY_BG, rounded=True)
    d.text(x + 0.22, 2.15, cw - 0.4, 0.8, [{"t": num, "size": 40, "bold": True, "color": BLUE}])
    d.text(x + 0.22, 3.0, cw - 0.4, 0.75, [{"t": head, "size": 19, "bold": True}])
    d.text(x + 0.22, 3.75, cw - 0.4, 1.7, [{"t": body, "size": 15, "color": INK2}])
d.text(0.7, 5.95, 12, 0.8, [{"t": "Every change is committed in git (branch **intern-work**) on top of an untouched copy of the "
                                  "original code, so each fix can be shown as a before / after.", "size": 16, "color": INK2}])

# 3 ---------------------------------------------------------------- how it works
d.new("How the system works", kicker="Background", notes="""
Here is how the system works. The Tobii eye tracker sits under the monitor. Because consumer Tobii devices don't give
raw data to our programs without a paid licence, the previous intern used a trick: the Tobii Ghost app draws a green
bubble where you look, and Python takes screenshots very fast and finds that green bubble. The bubble's position is the
gaze point. From the gaze points we detect fixations, which is where the eye holds still. The reader clicks a 'Yes,
confused' button whenever they are confused; those clicks are our ground truth labels. Finally a model learns to
predict confusion from the eye movements. Two consequences of the screenshot trick: the sampling rate is low, and we
can't measure pupil size, which is one of the best signals of confusion.""")
boxes = [("Tobii + Ghost", "draws a green bubble where you look"),
         ("Screenshots", "find the bubble → gaze point (x, y, time)"),
         ("Fixations", "eye holds ≥ 0.25 s within 25 px"),
         ("Labels", "reader clicks “Yes (Confused)”"),
         ("Model", "predicts confusion from eye movements")]
bw, ag = 2.05, 0.45
for i, (head, body) in enumerate(boxes):
    x = 0.7 + i * (bw + ag)
    d.rect(x, 2.2, bw, 2.1, BLUE_BG, rounded=True)
    d.text(x + 0.15, 2.4, bw - 0.3, 0.5, [{"t": head, "size": 19, "bold": True, "align": "c"}])
    d.text(x + 0.15, 2.95, bw - 0.3, 1.25, [{"t": body, "size": 15, "color": INK2, "align": "c"}])
    if i < len(boxes) - 1:
        d.arrow(x + bw + 0.07, 3.08, ag - 0.14, 0.34)
d.card(0.7, 4.8, 12.0, 1.65, "Limitation of the screenshot trick (no Tobii SDK)",
       ["Gaze is read off the screen, so we get 5–30 samples per second instead of the tracker's 60+, real saccades "
        "(20–80 ms) fall between samples, and pupil size — a strong confusion signal — is never measured."],
       fill=ORANGE_BG, body_size=16)

# 4 ---------------------------------------------------------------- key terms
d.new("Key terms used in this talk", kicker="Background", notes="""
A few terms I'll use. A fixation is when the eye stays still on one spot. A saccade is the quick jump between two
fixations. A regression is a jump backwards, which in reading means re-reading: either to the left on the same line or
up to an earlier line. An AOI, area of interest, is a box around one line of text. The confusion click is our ground
truth. For results I use ROC-AUC, where 0.5 means guessing and 1.0 is perfect, and PR-AUC, which focuses on finding the
rare confused moments and must be compared with its chance value. Leave-one-session-out means we always test on a
session, and therefore a text, that the model has never seen.""")
terms = [("Fixation", "the eye holds still on one spot (here: ≥ 0.25 s within 25 px)"),
         ("Saccade", "the fast jump between two fixations"),
         ("Regression", "a jump backwards — left on the same line or up to an earlier line (re-reading)"),
         ("AOI", "“area of interest”: a box around one text line, found by OCR"),
         ("Confusion click", "the reader presses “Yes (Confused)” — our ground truth"),
         ("ROC-AUC", "how well confused moments are ranked above normal ones: 0.5 = guessing, 1.0 = perfect"),
         ("PR-AUC", "how well the rare confused moments are found; compare with its chance value"),
         ("Leave-one-session-out", "train on all sessions but one, test on the unseen one, repeat for each")]
for i, (term, desc) in enumerate(terms):
    col, row = i % 2, i // 2
    x, y = 0.7 + col * 6.15, 1.85 + row * 1.2
    d.rect(x, y, 5.9, 1.05, GREY_BG, rounded=True)
    d.text(x + 0.2, y + 0.13, 5.5, 0.85, [{"t": term, "size": 17, "bold": True, "color": BLUE, "after": 1},
                                           {"t": desc, "size": 14, "color": INK2}])

# 5 ---------------------------------------------------------------- reproduction
d.new("Step 1 — I reproduced the Tobii pipeline end to end", kicker="Step 1 · Reproduce", notes="""
My first goal was to run the existing system exactly as documented, so I understood every step. I installed everything,
calibrated the Tobii, extracted the text lines with OCR, recorded myself reading, labelled the data and ran the live
prediction demo. I also installed and built the webcam app; only the camera test is pending because we don't have a
webcam on this PC. While doing this I found four problems the handover README didn't mention — on the right. The most
serious one is the first: the recording script appends to old files, so new data would silently mix with old data.""")
d.text(0.7, 1.85, 5.6, 0.45, [{"t": "What I ran on the lab PC", "size": 19, "bold": True}])
d.bullets(0.7, 2.4, 5.6, 4.3, [
    "Set up Python, Tesseract OCR, Tobii Experience and Ghost",
    "Verified the green-bubble detection: found in **10 / 10** test frames",
    "Extracted text-line boxes, recorded a session, labelled it",
    "Ran the live prediction demo with the shipped model",
    "Installed and built the webcam app (camera test pending — no webcam yet)"], size=17)
d.text(6.8, 1.85, 5.9, 0.45, [{"t": "Four traps the README didn't mention", "size": 19, "bold": True}])
traps = [("Recording appends to old files", "new sessions would mix with old data"),
         ("OCR screenshot is instant", "it captured the terminal, not the text"),
         ("label.py read the wrong file", "an old file, not the new session"),
         ("Shipped test file → crash", "19 text lines, but the model needs exactly 18")]
for i, (h, b) in enumerate(traps):
    d.card(6.8, 2.4 + i * 1.08, 5.9, 0.95, h, b, fill=ORANGE_BG, body_size=14, head_size=16)

# 6 ---------------------------------------------------------------- data audit
d.new("Step 2 — What data the previous intern left", kicker="Step 2 · Audit", notes="""
Next I checked what data exists. There are 20 recording sessions from July 2025, each on a different page of text, with
1,553 fixations. Importantly, every confusion click was saved with an exact timestamp, so I could relabel the data by
time without collecting anything new. Two weaknesses: the sampling rate was only 5 to 8 samples per second, and nobody
recorded who read the texts — the lab doesn't know whether it was one person or several. That limits what we can claim,
and it's why new data with known participants is the next step.""")
stats = [("20", "sessions (July 2025), a different text page in each", BLUE),
         ("1,553", "fixations — every one maps to exactly one session", BLUE),
         ("67", "confusion clicks, all with exact timestamps → relabelling by time is possible", BLUE),
         ("5–8 Hz", "sampling rate: only 2–3 samples per fixation", ORANGE),
         ("Unknown", "who read the texts — not recorded, nobody in the lab knows", ORANGE)]
sw = 2.25
for i, (num, cap, col) in enumerate(stats):
    d.stat(0.7 + i * (sw + 0.18), 1.95, sw, 2.6, num, cap, color=col)
d.card(0.7, 4.95, 12.0, 1.55, "Hidden problem: all 20 sessions were appended into one file",
       ["They all used the first session's column names, although each session had its own page and its own 18 line "
        "boxes. Three rows even had an extra column, so the file could not be read normally."],
       fill=ORANGE_BG, body_size=16)

# 7 ---------------------------------------------------------------- leakage
d.new("The old model learned where confusing lines were — not how confused eyes move", kicker="Step 2 · Audit", notes="""
This is the main problem with the old model. Its inputs included the screen position and one column per text line. But
because all sessions shared the same column names, the column 'AOI_5' simply means 'the fourth line on the screen' — a
different sentence in every text. On top of that, a click on a line labelled every fixation on that line as confused,
for the whole session. So the model could learn 'the fourth line tends to be confusing', which can't work on any new
text. The chart shows it measured: 72 percent of the model's decision comes from which line, and only 18 percent from
eye behaviour. In the live demo I saw the same thing: predictions only fired on certain lines and never off the text.""")
d.picture(os.path.join(CHARTS, "old_model_reliance.png"), 0.7, 2.0, w=6.3)
d.bullets(7.4, 1.95, 5.3, 4.7, [
    "Column “AOI_5” = **the 4th line on screen** — a different sentence in every text",
    "A click on a line marked **every** fixation on that line as confused, for the whole session",
    "So the model can learn “line 4 is confusing” — useless on any new text",
    "Seen live: “Confused” fired on only **8 of 18** lines, and **never** off the text (0 of 70)"], size=17, after=10)
d.card(0.7, 4.85, 6.3, 1.7, "Measured in Experiment 2",
       "Tested on unseen sessions instead of a random split, the old setup drops from AUC 0.87 to 0.74.",
       fill=ORANGE_BG, body_size=16)

# 8 ---------------------------------------------------------------- measurement bugs
d.new("Step 3 — Three measurement bugs I fixed", kicker="Step 3 · Fix", notes="""
Before running experiments I fixed how eye movements are measured, because the features are the model's inputs. First,
the regression flag: it fired on any upward movement, even one pixel. 89 percent of its flags were just the gaze
jittering on the same line, median 6 pixels, while it missed 115 of 180 real leftward re-reads. The new rule only counts
a jump left on the same line or up to an earlier line. Second, the time between fixations was measured from the start of
the previous fixation, so it included that whole fixation; I corrected it. Third, a single missed screenshot, for
example a blink, threw away the fixation in progress; now short dropouts are tolerated. All of this now lives in one
shared module, gaze_core.py, instead of six copied versions.""")
d.picture(os.path.join(CHARTS, "regression_flag.png"), 0.7, 1.95, w=6.2)
d.text(0.7, 4.15, 6.2, 0.9, [{"t": "89% of old flags were ≤ 25 px jitter on the same line (median 6 px); the old flag "
                                    "also missed 115 of 180 real leftward re-reads.", "size": 15, "color": INK2}])
d.card(7.3, 1.95, 5.4, 1.45, "1 · Regression flag", "Any upward move → now: left on the same line, or up to an earlier line")
d.card(7.3, 3.55, 5.4, 1.45, "2 · Time between fixations", "Was start-to-start (median 1.3 s) → now end-to-start (0.65 s)")
d.card(7.3, 5.15, 5.4, 1.45, "3 · Fixations lost on blinks", "One missed screenshot dropped the fixation → 68 fixations recovered")
d.card(0.7, 5.15, 6.2, 1.45, "One shared module: gaze_core.py",
       "The same code had been copy-pasted into six scripts; now analysis and recording use one version.",
       fill=BLUE_BG, head_color=BLUE)

# 9 ---------------------------------------------------------------- checks
d.new("How I made sure the fixes and results are correct", kicker="Step 3 · Fix", notes="""
Changing code is risky, so I checked every step. First, before applying any fix, I rewrote the original logic and
showed it reproduces 99.1 percent of the fixations the original script logged — so any difference afterwards comes from
the intended fixes. Second, a new recording processed live and offline gives identical results. Third, I checked the
results don't depend on my choice of window length, and that the effect isn't caused by looking at the click button.
These checks also caught two bugs in my own code — a time-zone offset and a missing-value issue — which I fixed and
reported; they changed the results by at most 0.01.""")
checks = [("99.1%", "Rewrite reproduces the original fixations (1,539 / 1,553) before any fix"),
          ("70 / 70", "A new recording gives identical fixations live and offline"),
          ("5 · 10 · 15 s", "Results hold for every window length; 10 s was fixed before seeing results"),
          ("0", "fixations on the click button inside the “confused” windows — not a button artefact")]
for i, (num, cap) in enumerate(checks):
    x = 0.7 + (i % 2) * 6.15
    y = 1.95 + (i // 2) * 1.75
    d.rect(x, y, 5.9, 1.55, GREY_BG, rounded=True)
    d.text(x + 0.25, y + 0.2, 2.2, 1.15, [{"t": num, "size": 30, "bold": True, "color": BLUE}], anchor="middle")
    d.text(x + 2.45, y + 0.2, 3.25, 1.15, [{"t": cap, "size": 15, "color": INK2}], anchor="middle")
d.card(0.7, 5.55, 12.0, 1.0, "These checks caught two bugs in my own code — fixed and reported",
       "a 5.5-hour time-zone offset, and the text “None” being read as a missing value; results moved by ≤ 0.01 AUC",
       fill=BLUE_BG, head_color=BLUE, body_size=15, head_size=16)

# 10 --------------------------------------------------------------- experiment 1 design
d.new("Experiment 1 — Label confusion by time, not by place", kicker="Step 4 · Experiments", notes="""
Experiment 1 changes how labels are made. The old rule asked: was this fixation ever on a line that was clicked? The new
rule asks: did this fixation happen in the seconds before a click? That describes when the reader was confused, and it
doesn't depend on line positions, so it can't leak position into the model. I tried windows of 5, 10 and 15 seconds.
With 15 seconds, 35 percent of fixations are labelled confused — about the same as the old rule — so the 'implausibly
high 36 percent' mentioned in the README comes mostly from the long window. Shorter windows give fewer, cleaner labels.""")
# timeline diagram
tl_y = 2.75
d.rect(0.9, tl_y, 11.4, 0.06, GREY)
d.rect(6.2, tl_y - 0.55, 3.6, 1.15, BLUE_BG, rounded=True)
d.rect(9.8, tl_y - 0.75, 0.07, 1.55, ORANGE)
d.text(9.97, tl_y - 0.8, 2.6, 0.5, [{"t": "click: “Yes (Confused)”", "size": 15, "bold": True, "color": ORANGE}])
d.text(6.2, tl_y + 0.7, 3.6, 0.4, [{"t": "5 / 10 / 15 s before the click → confused", "size": 15, "bold": True, "color": BLUE, "align": "c"}])
for xx in [1.3, 1.9, 2.6, 3.2, 3.9, 4.6, 5.2, 5.8, 6.5, 7.0, 7.5, 8.1, 8.6, 9.2, 10.4, 11.0, 11.7]:
    col = BLUE if 6.2 <= xx <= 9.8 else GREY
    d.rect(xx, tl_y - 0.12, 0.3, 0.3, col, rounded=True)
d.text(0.9, tl_y + 0.7, 5, 0.4, [{"t": "fixations over time →", "size": 14, "color": MUTED}])
d.card(0.7, 4.05, 5.9, 1.25, "Old rule (by place)", "Fixation on a clicked line, at any time in the session → confused",
       fill=ORANGE_BG)
d.card(0.7, 5.45, 5.9, 1.25, "New rule (by time)", "Fixation in the seconds before a click → confused",
       fill=BLUE_BG, head_color=BLUE)
rows = [("Window", "Labelled confused"), ("5 s", "16%"), ("10 s", "26%"), ("15 s", "35%"), ("Old rule", "34%")]
for i, (a, b) in enumerate(rows):
    y = 4.05 + i * 0.53
    if i == 0:
        d.rect(7.1, y, 5.6, 0.5, GREY_BG)
    d.text(7.3, y + 0.08, 2.5, 0.4, [{"t": a, "size": 16, "bold": i == 0}])
    d.text(9.9, y + 0.08, 2.6, 0.4, [{"t": b, "size": 16, "bold": i == 0, "align": "r"}])

# 11 --------------------------------------------------------------- experiment 1 result
d.new("Before a confusion click, the eyes really do behave differently", kicker="Step 4 · Experiments", notes="""
With time-based labels I compared eye behaviour in the 5 seconds before each click with normal reading, inside each
session, so every session is compared with itself. The chart shows the regression rate: every line is one session, and
every single line goes up. The median rises from 11 to 28 percent. Fixations also get longer, and jumps get shorter.
All three are statistically significant. These are exactly the signs of reading difficulty described in the literature,
so the gaze data contains a real confusion signal. A caveat: this is a difference in averages; whether it allows
prediction is tested in Experiment 2.""")
d.picture(os.path.join(CHARTS, "before_click.png"), 0.7, 1.75, h=4.95)
items = [("19 / 19", "sessions: more re-reading before a click (median 11% → 28%)"),
         ("17 / 19", "sessions: longer fixations (+0.13 s)"),
         ("19 / 19", "sessions: shorter jumps between fixations (−95 px)")]
for i, (num, cap) in enumerate(items):
    y = 1.95 + i * 1.45
    d.rect(7.3, y, 5.4, 1.3, GREY_BG, rounded=True)
    d.text(7.5, y + 0.15, 1.9, 1.0, [{"t": num, "size": 28, "bold": True, "color": BLUE}], anchor="middle")
    d.text(9.45, y + 0.15, 3.1, 1.0, [{"t": cap, "size": 15, "color": INK2}], anchor="middle")
d.text(7.3, 6.35, 5.4, 0.5, [{"t": "Each p < 0.001 (paired test per session). Classic signs of reading difficulty.",
                              "size": 14, "color": INK2}])

# 12 --------------------------------------------------------------- experiment 2 design
d.new("Experiment 2 — An honest test of whether confusion can be predicted", kicker="Step 4 · Experiments", notes="""
Experiment 2 asks whether we can predict confusion, measured honestly. Instead of single fixations, I use 10-second
windows of reading, because confusion lasts a few seconds. A window counts as confused if it ends 0 to 3 seconds before
a click — it contains the reading that led to the click, but not the click itself. The features describe only how the
eyes move: regression rates, fixation durations, jump sizes and so on — no position and no line identity. The test is
leave-one-session-out: train on 19 sessions, test on the 20th, and repeat for all. Since each session is a different
text, this measures performance on new text. Confused windows are rare, 7 percent, which matters for reading PR-AUC.""")
design = [("Unit", "10-second windows of reading, moved forward 1 s at a time"),
          ("Label", "“Confused” if the window ends 0–3 s before a click; windows containing a click are dropped"),
          ("Features", "13 behaviour measures — re-reading, fixation duration, jump size… No position, no line identity"),
          ("Test", "Leave-one-session-out: train on 19 sessions, test on the unseen 20th; every session is a new text"),
          ("Models", "Random Forest and XGBoost, fixed random seed; code and results committed (train.py)"),
          ("Data", "2,357 windows from 20 sessions; 166 confused (7%)")]
for i, (h, b) in enumerate(design):
    x = 0.7 + (i % 3) * 4.1
    y = 1.95 + (i // 3) * 2.35
    d.card(x, y, 3.85, 2.15, h, b, fill=BLUE_BG if i in (2, 3) else GREY_BG,
           head_color=BLUE if i in (2, 3) else INK, body_size=16, head_size=19)

# 13 --------------------------------------------------------------- experiment 2 results
d.new("Eye behaviour alone detects confusion on unseen texts: AUC 0.76", kicker="Step 4 · Experiments", notes="""
Here are the results. The blue bar is the new model: behaviour only, tested on unseen sessions, AUC 0.76. In PR-AUC terms
it finds confused windows about four times better than chance, 0.29 versus 0.07, and it works in 18 of the 19 sessions
with clicks. The orange bars are the old setup: 0.87 on a random split, but only 0.74 when tested on unseen sessions.
That drop of 0.13 is the over-optimism caused by the leakage, now measured. Note the old and new numbers aren't directly
comparable because the labels differ; the key point is that the old 0.87 was inflated. The grey bar shows a single
feature, regression rate, already gets 0.71, so machine learning adds a modest improvement.""")
d.picture(os.path.join(CHARTS, "exp2_auc.png"), 0.6, 1.95, w=7.4)
res = [("0.29", "PR-AUC vs 0.07 by chance — about 4× better at finding confused windows"),
       ("18 / 19", "sessions with clicks where the model beats guessing"),
       ("−0.13", "AUC the old setup loses on unseen sessions → the leakage, measured")]
for i, (num, cap) in enumerate(res):
    y = 1.95 + i * 1.45
    d.rect(8.35, y, 4.35, 1.3, GREY_BG, rounded=True)
    d.text(8.55, y + 0.15, 1.6, 1.0, [{"t": num, "size": 26, "bold": True, "color": BLUE if i < 2 else ORANGE}], anchor="middle")
    d.text(10.2, y + 0.15, 2.4, 1.0, [{"t": cap, "size": 14, "color": INK2}], anchor="middle")
d.text(0.7, 5.15, 7.3, 1.5, [{"t": "Old and new numbers are not directly comparable (different labels and units). "
                                    "What matters: the old 0.87 came from a random split and was inflated; the new "
                                    "0.76 is measured on unseen texts from the start.", "size": 15, "color": INK2}])

# 14 --------------------------------------------------------------- meaning
d.new("What the results mean — and what they don't yet", kicker="Step 4 · Experiments", notes="""
It's important to be clear about what we can and cannot claim. We can say that eye behaviour carries a real confusion
signal, that it works on texts the model has never seen, and that the old accuracy was inflated. We cannot yet say it
works across people, because nobody knows who read the old texts — possibly only one person. My own session on the new
setup was too small to judge: 46 windows with 9 confused, and it was recorded at a much higher sampling rate, which
changes the features. And with 5 to 8 samples per second and no pupil data, we can't study fast saccades or pupil
dilation. These open points define the next experiments.""")
d.rect(0.7, 1.95, 5.9, 3.95, BLUE_BG, rounded=True)
d.text(0.95, 2.15, 5.4, 0.5, [{"t": "We can say", "size": 21, "bold": True, "color": BLUE}])
d.bullets(0.95, 2.8, 5.4, 3.0, [
    "Eye behaviour carries a real confusion signal",
    "It works on texts the model has never seen (AUC 0.76)",
    "The old 0.87 was inflated by leakage"], size=18, after=12)
d.rect(6.85, 1.95, 5.85, 3.95, ORANGE_BG, rounded=True)
d.text(7.1, 2.15, 5.4, 0.5, [{"t": "We cannot say yet", "size": 21, "bold": True, "color": ORANGE}])
d.text(7.1, 2.8, 5.4, 3.0, [
    {"t": "That it works **across people** — the old readers are unknown", "size": 18, "bullet": True, "bcolor": ORANGE, "after": 12},
    {"t": "That it works on the **new setup** — my test session was too small (46 windows, 9 confused) and recorded at 3× the sampling rate",
     "size": 18, "bullet": True, "bcolor": ORANGE, "after": 12},
    {"t": "Anything about fast saccades or **pupil size** — not measurable at 5–8 Hz", "size": 18, "bullet": True,
     "bcolor": ORANGE, "after": 12}])
d.text(0.7, 6.15, 12.0, 0.5, [{"t": "The open points on the right are exactly what Experiments 3 and 4 are designed to answer.",
                               "size": 17, "bold": True, "color": INK2}])

# 15 --------------------------------------------------------------- recording fixed
d.new("Step 5 — The recording script is fixed for new data", kicker="Step 5 · Ready for new data", notes="""
Before collecting new data, I fixed the recording script, so new sessions don't repeat the old problems. The new script,
collect.py, saves each session in its own folder named by participant, text and time, and never appends. It captures
the text lines at the start of each session, uses the same validated code as the analysis, and records the setup in a
meta file. Removing unnecessary pauses raised the sampling rate from 18 to 29 samples per second on our PC — four to
five times the old data. I tested it live with the Tobii. I also wrote a quality check that runs after each recording,
and a step-by-step protocol for running participants.""")
d.picture(os.path.join(CHARTS, "sampling_rate.png"), 0.7, 2.0, w=5.9)
d.text(0.7, 4.45, 5.9, 0.9, [{"t": "Same tracker, same PC — only the code changed. Tested live with the Tobii.",
                                    "size": 15, "color": INK2}])
d.bullets(7.0, 1.95, 5.7, 4.8, [
    "**collect.py** replaces cnn.py: one folder per session (participant_text_time), never appends",
    "Text lines captured inside each session; same validated code as the analysis",
    "meta.json records sampling rate, screen, settings and code version",
    "**check_session.py**: PASS / WARN / FAIL quality check right after each recording",
    "**PROTOCOL.md**: step-by-step procedure and word-for-word participant instructions"], size=16, after=10)

# 16 --------------------------------------------------------------- webcam
d.new("Webcam track — set up, camera test pending", kicker="Parallel track", notes="""
The second intern built a webcam version that runs in the browser without an eye tracker. I installed it and it builds
and runs; the only missing part is a test with a real camera, because this PC has none. It's important to know its
limits: a webcam's gaze error is around 150 pixels, while a text line is about 30 pixels tall, so it can only tell which
paragraph you look at, not which line. It can't measure pupil size either. On the other hand it can see blinks and brow
furrowing, which the Tobii can't. The plan is to add logging and a confused button, add calibration, and then record the
webcam and the Tobii at the same time to measure exactly how much accuracy a webcam loses.""")
d.card(0.7, 1.95, 3.85, 3.3, "Status", ["Installed, builds and runs in the browser",
                                       "Camera test pending — no webcam on this PC",
                                       "Today: no logging, no calibration, gaze follows the head more than the eyes"],
       body_size=16, head_size=19)
d.card(4.75, 1.95, 3.85, 3.3, "What a webcam can do", ["Fixations and regressions (coarse)",
                                                        "Paragraph-level areas (not single lines)",
                                                        "Bonus: blink rate and brow furrow — the Tobii can't see these"],
       fill=BLUE_BG, head_color=BLUE, body_size=16, head_size=19)
d.card(8.8, 1.95, 3.9, 3.3, "What it cannot do", ["Line-level analysis: error ~150 px vs a ~30 px line",
                                                   "Pupil size",
                                                   "Saccade speed (only ~30 frames/s)"],
       fill=ORANGE_BG, head_color=ORANGE, body_size=16, head_size=19)
d.card(0.7, 5.4, 12.0, 1.3, "Plan (tasks W1 → W2 → W5)",
       "Add logging and a “confused” button → add calibration → record webcam and Tobii together to measure "
       "how much accuracy a webcam loses.", fill=BLUE_BG, head_color=BLUE, body_size=16)

# 17 --------------------------------------------------------------- next steps
d.new("Next steps — and what I need from you", kicker="Plan", notes="""
The next step is Experiment 3: new data with known participants. I propose three texts — easy, medium and hard — and 10
to 12 participants, each reading all three in a rotated order. The protocol and quality check are ready. With that data
I can test, for the first time, whether the result holds across people, by leaving one participant out at a time. For
this I need four decisions from you: approval of the texts and participant plan, the ethics and consent process, whether
we can get a Tobii Pro device or SDK licence for pupil size and real saccades, and a webcam for the webcam track.""")
d.text(0.7, 1.85, 6.0, 0.5, [{"t": "Experiment 3 — new data, known readers", "size": 20, "bold": True}])
d.bullets(0.7, 2.45, 6.0, 4.2, [
    "3 texts: easy, medium, hard (hard texts produce real confusion)",
    "10–12 participants (anonymous ids P01, P02, …), all three texts each, in rotated order",
    "Protocol, recording script and quality check are ready",
    "Then: **leave-one-participant-out** — does it work across people?",
    "Later: record webcam and Tobii together to measure the webcam's accuracy"], size=17, after=10)
d.rect(7.1, 1.85, 5.6, 4.8, BLUE_BG, rounded=True)
d.text(7.35, 2.05, 5.2, 0.5, [{"t": "Decisions I need", "size": 20, "bold": True, "color": BLUE}])
decisions = ["Approve the texts and the 10–12 participant plan",
             "Ethics / consent process for recording participants",
             "Tobii Pro device or SDK licence? (pupil size + real saccades — Experiment 4)",
             "A webcam for the webcam track"]
d.text(7.35, 2.7, 5.15, 3.8, [{"t": f"{i}.  {t}", "size": 17, "after": 12} for i, t in enumerate(decisions, 1)])

# 18 --------------------------------------------------------------- summary
d.new("Summary", kicker="Summary", notes="""
To summarise: the pipeline runs end to end, and every change is versioned and checked. The old model's accuracy was not
valid, because it learned where the confusing lines were rather than how confused eyes move. After fixing the
measurements and evaluating honestly, eye behaviour alone detects the build-up to a confusion click on unseen texts with
an AUC of 0.76. The next step is new data with known readers, to test whether this holds across people. Thank you — I'm
happy to take questions.""")
takeaways = [("1", "The pipeline runs end to end; every fix is versioned and checked."),
             ("2", "The old model learned where confusing lines were — its accuracy was not valid."),
             ("3", "Eye behaviour alone detects confusion on unseen texts: AUC 0.76, ~4× chance PR-AUC."),
             ("4", "Next: 10–12 known readers to test whether this holds across people.")]
for i, (n, t) in enumerate(takeaways):
    y = 1.9 + i * 1.12
    d.rect(0.7, y, 12.0, 1.02, BLUE_BG if i == 2 else GREY_BG, rounded=True)
    d.text(0.95, y + 0.12, 0.6, 0.8, [{"t": n, "size": 30, "bold": True, "color": BLUE}], anchor="middle")
    d.text(1.7, y + 0.12, 10.8, 0.8, [{"t": t, "size": 20, "bold": i == 2}], anchor="middle")
d.text(0.7, 6.45, 12, 0.35, [{"t": "Thank you — questions?", "size": 16, "bold": True, "color": BLUE}])

# 19-20 ------------------------------------------------------------ Q&A
qa = [
    ("Why is the new accuracy lower than the old one?",
     "The old 0.87 came from a random split: the model was tested on other fixations from the same texts and lines it "
     "trained on. Tested on unseen sessions it falls to 0.74. Our 0.76 is measured on unseen texts from the start."),
    ("Why 10-second windows instead of single fixations?",
     "Confusion lasts seconds, single fixations at 5–8 Hz are noisy, and the literature uses 10–15 s windows. "
     "Results hold for 5, 10 and 15 s."),
    ("Why “0–3 s before the click”?",
     "Readers click after they notice they're confused. The window holds the reading that led to the click, "
     "but not the movement to the button and back."),
    ("Isn't the effect just looking at the button?",
     "No: zero fixations landed on the button inside the windows, and the result is the same using only fixations "
     "on text lines."),
    ("Why does your own session (s21) score near chance?",
     "Only 46 windows (9 confused), a different reader, and 3× the sampling rate, which changes features such as "
     "fixation counts. Too small to judge — that's why Experiment 3 is needed."),
    ("Why not use the Tobii SDK directly?",
     "Consumer Tobii devices block raw-data access without a licence. A Tobii Pro device or SDK licence would give "
     "60+ Hz and pupil size (Experiment 4)."),
    ("How do you know your fixes didn't break anything?",
     "The rewrite reproduces 99.1% of the original fixations, live and offline processing match exactly, and all "
     "results were re-run after every fix."),
    ("What does PR-AUC 0.29 mean?",
     "Only 7% of windows are confused, so a random guess scores 0.07. 0.29 means confused windows are found about "
     "4× better than chance."),
]
for page in range(2):
    d.new(f"Questions I may be asked ({page + 1}/2)", kicker="Appendix · preparation", notes="""
Preparation slide: likely questions with short answers. Keep it in the appendix and show it only if asked.""")
    for i, (q, a) in enumerate(qa[page * 4:(page + 1) * 4]):
        x = 0.7 + (i % 2) * 6.15
        y = 1.9 + (i // 2) * 2.45
        d.card(x, y, 5.9, 2.25, q, a, body_size=14, head_size=16)

# 21 --------------------------------------------------------------- where things are
d.new("Where everything is (branch intern-work)", kicker="Appendix · reference", notes="""
Reference slide: where each piece of work lives in the repository, in case you want to look at the code or results.""")
files = [("README.md §10.9", "all findings, with status: fixed / open"),
         ("gaze_core.py", "shared fixation code with the three fixes"),
         ("validate_gaze_core.py", "check: rewrite reproduces 99.1% of the original"),
         ("rebuild_fixations.py", "corrected fixations, one file per session"),
         ("label.py", "Experiment 1 — labels by time window"),
         ("train.py", "Experiment 2 — windows, features, leave-one-session-out"),
         ("processed/exp2_results.json", "all Experiment 2 numbers"),
         ("collect.py", "new recording script (replaces cnn.py)"),
         ("check_session.py", "quality check after each recording"),
         ("PROTOCOL.md", "procedure for Experiment 3 sessions")]
for i, (f, desc) in enumerate(files):
    col, row = i % 2, i // 2
    x, y = 0.7 + col * 6.15, 1.9 + row * 0.95
    d.rect(x, y, 5.9, 0.82, GREY_BG, rounded=True)
    d.text(x + 0.2, y + 0.1, 5.5, 0.65, [{"t": f, "size": 15, "bold": True, "after": 0},
                                           {"t": desc, "size": 13, "color": INK2}])

out = os.path.join(HERE, "Gaze_Confusion_Progress.pptx")
d.save(out)
print(f"{d.n} slides -> {out}")
print("overflow:", *(d.overflows or ["none"]), sep="\n  ")
